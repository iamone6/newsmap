# /// script
# requires-python = ">=3.11"
# dependencies = ["psycopg[binary]>=3.2", "openpyxl>=3.1", "python-dotenv>=1.0"]
# ///
"""행정안전부 행정동 코드(KIKcd_H) 파일을 regions 테이블에 적재한다.

입력: jumin.mois.go.kr 에서 받은 KIKcd_H.YYYYMMDD.xlsx, 또는 같은 컬럼의 CSV(UTF-8 / CP949)
  컬럼: 행정동코드, 시도명, 시군구명, 읍면동명, 생성일자, 말소일자

동작:
  - code 기준 upsert. 기존 행은 지우지 않는다 (폐지 코드는 말소일자 → valid_to).
  - level 은 코드 모양으로 판단: XX00000000 → 시도, XXXXX00000 → 시군구, 그 외 → 읍면동
  - parent_code: 읍면동 → 소속 시군구, 일반구(예: 수원시 영통구) → 시(수원시), 시군구 → 시도
  - 전체를 한 트랜잭션으로 처리한다.

사용:
  uv run apps/rag/scripts/load_regions.py KIKcd_H.20260101.xlsx
  python apps/rag/scripts/load_regions.py regions.csv --env apps/rag/.env --dry-run
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import os
import sys
from pathlib import Path

import psycopg
from dotenv import load_dotenv

COLUMNS = ["행정동코드", "시도명", "시군구명", "읍면동명", "생성일자", "말소일자"]


def read_rows(path: Path) -> list[dict]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        sheet = load_workbook(path, read_only=True, data_only=True).active
        it = sheet.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(it)]
        return [dict(zip(header, r)) for r in it]

    for encoding in ("utf-8-sig", "cp949"):
        try:
            with path.open(encoding=encoding, newline="") as f:
                return list(csv.DictReader(f))
        except UnicodeDecodeError:
            continue
    raise SystemExit(f"CSV 인코딩을 읽을 수 없습니다 (UTF-8 / CP949): {path}")


def clean(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def parse_date(value) -> dt.date | None:
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    digits = "".join(ch for ch in str(value) if ch.isdigit())
    if len(digits) >= 8:
        return dt.date(int(digits[:4]), int(digits[4:6]), int(digits[6:8]))
    raise ValueError(f"날짜 형식을 알 수 없습니다: {value!r}")


def to_region(row: dict) -> tuple | None:
    code = clean(row.get("행정동코드"))
    if code is None:
        return None
    code = code.split(".")[0].zfill(10)  # 엑셀/CSV 에서 숫자로 읽힌 경우 보정
    if len(code) != 10 or not code.isdigit():
        raise ValueError(f"행정동코드 형식 오류: {code!r}")

    sido, sigungu, emd = clean(row.get("시도명")), clean(row.get("시군구명")), clean(row.get("읍면동명"))
    if code[2:] == "00000000":
        level, sigungu, emd = 1, None, None
    elif code[5:] == "00000":
        level, emd = 2, None
    else:
        level = 3

    full_name = " ".join(n for n in (sido, sigungu, emd) if n)
    return (
        code,
        level,
        sido,
        sigungu,
        emd,
        full_name,
        parse_date(row.get("생성일자")),
        parse_date(row.get("말소일자")),
    )


def main() -> int:
    parser = argparse.ArgumentParser(description="행정동 코드(KIKcd_H)를 regions 테이블에 적재")
    parser.add_argument("file", type=Path, help="KIKcd_H.YYYYMMDD.xlsx 또는 CSV")
    parser.add_argument("--env", type=Path, default=Path(__file__).resolve().parents[1] / ".env",
                        help="DB_* 값을 읽을 .env (기본: apps/rag/.env)")
    parser.add_argument("--dry-run", action="store_true", help="적재 후 롤백")
    args = parser.parse_args()

    if args.env.exists():
        load_dotenv(args.env)

    rows = read_rows(args.file)
    missing = [c for c in COLUMNS if rows and c not in rows[0]]
    if missing:
        raise SystemExit(f"필수 컬럼이 없습니다: {missing}")

    regions = {}
    for row in rows:
        region = to_region(row)
        if region:
            regions[region[0]] = region  # 같은 코드가 여러 번 나오면 마지막 행 사용
    if not regions:
        raise SystemExit("적재할 행이 없습니다.")

    conninfo = dict(
        host=os.environ.get("DB_HOST", "127.0.0.1"),
        port=os.environ.get("DB_PORT", "5432"),
        dbname=os.environ["DB_DATABASE"],
        user=os.environ["DB_USERNAME"],
        password=os.environ["DB_PASSWORD"],
    )

    with psycopg.connect(**conninfo) as conn, conn.cursor() as cur:
        cur.execute("""
            CREATE TEMP TABLE regions_import (
                code varchar(10) PRIMARY KEY, level smallint, sido_name text, sigungu_name text,
                emd_name text, full_name text, valid_from date, valid_to date
            ) ON COMMIT DROP
        """)
        with cur.copy("COPY regions_import FROM STDIN") as copy:
            for region in regions.values():
                copy.write_row(region)

        # 1) upsert (parent_code 는 2단계에서 채운다: FK 순서 문제 회피)
        cur.execute("""
            INSERT INTO regions (code, level, sido_name, sigungu_name, emd_name, full_name, valid_from, valid_to)
            SELECT code, level, sido_name, sigungu_name, emd_name, full_name, valid_from, valid_to
            FROM regions_import
            ON CONFLICT (code) DO UPDATE SET
                level = EXCLUDED.level, sido_name = EXCLUDED.sido_name,
                sigungu_name = EXCLUDED.sigungu_name, emd_name = EXCLUDED.emd_name,
                full_name = EXCLUDED.full_name, valid_from = EXCLUDED.valid_from,
                valid_to = EXCLUDED.valid_to
            RETURNING (xmax = 0) AS inserted
        """)
        results = cur.fetchall()
        inserted = sum(1 for (flag,) in results if flag)

        # 2) parent_code
        #   읍면동 → 같은 앞 5자리의 시군구
        #   일반구("수원시 영통구") → 같은 시도의 시("수원시"), 없으면 시도
        #   시군구 → 시도
        cur.execute("""
            WITH parents AS (
                SELECT i.code,
                       CASE i.level
                           WHEN 3 THEN substr(i.code, 1, 5) || '00000'
                           WHEN 2 THEN coalesce(
                               (SELECT p.code FROM regions p
                                 WHERE p.level = 2
                                   AND p.sido_name = i.sido_name
                                   AND position(' ' IN i.sigungu_name) > 0
                                   AND p.sigungu_name = split_part(i.sigungu_name, ' ', 1)
                                   AND p.code <> i.code
                                 ORDER BY p.valid_to NULLS FIRST, p.code
                                 LIMIT 1),
                               substr(i.code, 1, 2) || '00000000')
                       END AS parent_code
                FROM regions_import i
                WHERE i.level > 1
            )
            UPDATE regions r
            SET parent_code = p.parent_code
            FROM parents p
            WHERE r.code = p.code
              AND EXISTS (SELECT 1 FROM regions x WHERE x.code = p.parent_code)
              AND r.parent_code IS DISTINCT FROM p.parent_code
        """)
        cur.execute("""
            SELECT count(*) FILTER (WHERE r.level > 1 AND r.parent_code IS NULL),
                   count(*) FILTER (WHERE NOT EXISTS (SELECT 1 FROM regions_import i WHERE i.code = r.code))
            FROM regions r
        """)
        orphan, not_in_file = cur.fetchone()

        levels = {1: 0, 2: 0, 3: 0}
        for region in regions.values():
            levels[region[1]] += 1
        active = sum(1 for region in regions.values() if region[7] is None)

        print(f"파일: {args.file.name}  행 {len(regions)}개 (시도 {levels[1]}, 시군구 {levels[2]}, 읍면동 {levels[3]}, 현재 유효 {active})")
        print(f"신규 {inserted}개, 갱신 {len(results) - inserted}개")
        if orphan:
            print(f"경고: 상위 코드를 찾지 못한 행 {orphan}개 (parent_code NULL)")
        if not_in_file:
            print(f"참고: 이번 파일에 없는 기존 코드 {not_in_file}개 (삭제하지 않음)")

        if args.dry_run:
            conn.rollback()
            print("--dry-run: 롤백했습니다.")
        else:
            conn.commit()
    return 0


if __name__ == "__main__":
    sys.exit(main())
